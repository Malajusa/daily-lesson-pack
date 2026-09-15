from __future__ import annotations

import copy
import json
import tempfile
import unicodedata
from pathlib import Path

from pptx import Presentation

from render_provenance import validate_structural_text


def _normal(text: str) -> str:
    return " ".join(unicodedata.normalize("NFKC", text).split())


def _structural_records(render_manifest: dict) -> dict[tuple[int, int], dict]:
    if render_manifest.get("schema_version") != 1 or render_manifest.get("renderer") != "daily-lesson-pack":
        raise ValueError("Unsupported render manifest")
    result: dict[tuple[int, int], dict] = {}
    slides = render_manifest.get("slides")
    if not isinstance(slides, list):
        raise ValueError("Render manifest slides must be a list")
    for slide in slides:
        page = slide.get("slide_number")
        entries = slide.get("structural_text", [])
        errors = validate_structural_text(entries)
        if errors:
            raise ValueError("; ".join(errors))
        for entry in entries:
            key = (page, entry["shape_id"])
            if key in result:
                raise ValueError(f"Duplicate structural text binding {key}")
            result[key] = entry
    return result


def sanitise_structural_text(
    deck_path: Path,
    render_manifest: dict,
    out_path: Path,
    *,
    bound_shapes: set[tuple[int, int]] | None = None,
) -> list[str]:
    records = _structural_records(render_manifest)
    presentation = Presentation(deck_path)
    errors: list[str] = []
    bound_shapes = bound_shapes or set()
    seen_structural: set[tuple[int, int]] = set()
    for page, slide in enumerate(presentation.slides, 1):
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False) or not shape.text.strip():
                continue
            key = (page, shape.shape_id)
            entry = records.get(key)
            if entry is not None:
                seen_structural.add(key)
                if shape.name != entry["shape_name"] or _normal(shape.text) != _normal(entry["text"]):
                    errors.append(f"Structural renderer text mismatch at slide {page} shape {shape.shape_id}")
                    continue
                shape.text_frame.clear()
                continue
            if key not in bound_shapes:
                errors.append(f"Unrecognised unbound renderer text at slide {page} shape {shape.shape_id}")
    missing = set(records) - seen_structural
    if missing:
        errors.append("Render manifest structural text references missing shapes: " + ", ".join(map(str, sorted(missing))))
    if not errors:
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        presentation.save(out_path)
    return errors


def audit_repo_rendered_pack(manifest_path, content_path, context_path, render_manifest_path):
    from pack_evidence import audit_pack, digest, pages, read

    manifest_path = Path(manifest_path)
    render_manifest_path = Path(render_manifest_path)
    manifest = read(manifest_path)
    render_manifest = read(render_manifest_path)
    deck_entries = [artifact for artifact in manifest.get("artifacts", []) if artifact.get("role") == "deck"]
    if len(deck_entries) != 1:
        return ["Repository-rendered pack requires exactly one deck artifact"], read(content_path), manifest
    deck = deck_entries[0]
    deck_path = (manifest_path.parent / deck["path"]).resolve()
    bound = {
        (binding["page"], binding["shape_id"])
        for binding in manifest.get("bindings", [])
        if binding.get("artifact") == deck["id"] and "shape_id" in binding
    }
    with tempfile.TemporaryDirectory(prefix="dlp-render-audit-", dir=manifest_path.parent) as folder:
        folder = Path(folder)
        sanitised = folder / "deck-sanitised.pptx"
        errors = sanitise_structural_text(deck_path, render_manifest, sanitised, bound_shapes=bound)
        if errors:
            return errors, read(content_path), manifest
        audit_manifest = copy.deepcopy(manifest)
        for artifact in audit_manifest["artifacts"]:
            original = (manifest_path.parent / artifact["path"]).resolve()
            artifact["path"] = str(sanitised if artifact["id"] == deck["id"] else original)
            if artifact["id"] == deck["id"]:
                artifact["sha256"] = digest(sanitised)
                artifact["pages"] = pages(sanitised)
        temp_manifest = folder / "manifest.json"
        temp_manifest.write_text(json.dumps(audit_manifest), encoding="utf-8")
        core_errors, content, _ = audit_pack(temp_manifest, content_path, context_path)
        return core_errors, content, manifest
