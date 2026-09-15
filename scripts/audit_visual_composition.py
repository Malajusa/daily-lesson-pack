#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from pptx import Presentation


DATE_META_RE = re.compile(
    r"\b(?:Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday)\b.*\b20\d{2}\b|\bTerm\s+\d+\b.*\bWeek\s+\d+\b",
    re.I,
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _read(path: Path) -> dict:
    data = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError(f"{path} must contain a JSON object")
    return data


def _shape_map(slide):
    return {shape.shape_id: shape for shape in slide.shapes}


def _font_sizes(shape):
    sizes = []
    if getattr(shape, "has_text_frame", False):
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.size:
                    sizes.append(run.font.size.pt)
    return sizes


def _major_geometry(slide):
    rows = []
    for shape in slide.shapes:
        name = getattr(shape, "name", "") or ""
        if name.startswith(("DLP:panel:", "DLP:group:", "DLP:check-panel:")):
            rows.append(
                (
                    round(shape.left / 914400, 1),
                    round(shape.top / 914400, 1),
                    round(shape.width / 914400, 1),
                    round(shape.height / 914400, 1),
                )
            )
    return tuple(sorted(rows))


def _warning(code: str, record: dict, observation: str, source_blocks=None) -> dict:
    return {
        "code": code,
        "slide": record.get("slide_number"),
        "layout_id": record.get("layout_id"),
        "source_blocks": source_blocks if source_blocks is not None else record.get("source_blocks", []),
        "observation": observation,
    }


def audit_visual_composition(deck_path: Path, render_manifest_path: Path) -> dict:
    deck_path = Path(deck_path)
    render_manifest_path = Path(render_manifest_path)
    manifest = _read(render_manifest_path)
    errors: list[str] = []
    warnings: list[dict] = []
    if manifest.get("schema_version") != 1:
        errors.append("render manifest schema_version must be 1")
    if manifest.get("renderer") != "daily-lesson-pack":
        errors.append("render manifest renderer must be daily-lesson-pack")
    if manifest.get("deck_sha256") and manifest.get("deck_sha256") != sha256(deck_path):
        errors.append("render manifest deck_sha256 does not match deck")

    presentation = Presentation(deck_path)
    records = manifest.get("slides")
    if not isinstance(records, list) or len(records) != len(presentation.slides):
        errors.append("render manifest slide count does not match deck")
        records = []

    geometry = []
    for slide, record in zip(presentation.slides, records):
        shapes = _shape_map(slide)
        role = record.get("role", "")
        slide_number = int(record.get("slide_number") or 0)
        bindings = record.get("shape_bindings") or []
        geometry.append((_major_geometry(slide), role, record.get("layout_id"), slide_number))

        for binding in bindings:
            shape = shapes.get(binding.get("shape_id"))
            if shape is None:
                warnings.append(
                    _warning(
                        "COMPOSITION.MISSING_BINDING",
                        record,
                        "Render manifest references a missing shape.",
                        [binding.get("source_block_id")],
                    )
                )
                continue
            field = str(binding.get("source_field", ""))
            if field == "teacher_note" or str(binding.get("source_block_id", "")).endswith(":teacher_note"):
                warnings.append(
                    _warning(
                        "COMPOSITION.TEACHER_CONTENT",
                        record,
                        "Teacher-only content is bound to a projected shape.",
                        [binding.get("source_block_id")],
                    )
                )
            if field == "paragraph" and shape.width / 914400 > 8.2:
                warnings.append(
                    _warning(
                        "COMPOSITION.PARAGRAPH_MEASURE",
                        record,
                        f"Paragraph measure is {shape.width / 914400:.2f} in; preferred maximum is 8.2 in.",
                        [binding.get("source_block_id")],
                    )
                )

        if slide_number != 1 and role != "transition":
            texts = [
                shape.text.strip()
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False) and shape.text.strip()
            ]
            texts.extend(record.get("metadata_labels") or [])
            if any(DATE_META_RE.search(text) for text in texts):
                warnings.append(
                    _warning(
                        "COMPOSITION.REPEATED_METADATA",
                        record,
                        "Date/term/week metadata appears on an ordinary projected slide.",
                    )
                )

        if role == "morning_work" and len(record.get("groups") or {}) < 2:
            warnings.append(
                _warning(
                    "COMPOSITION.MORNING_WORK_GROUPING",
                    record,
                    "Morning Work does not record two broad working areas.",
                )
            )

        if role == "success_criteria":
            bound_shapes = []
            for binding in bindings:
                shape = shapes.get(binding.get("shape_id"))
                if shape is not None and str(binding.get("source_field")) == "success_criteria":
                    bound_shapes.append(shape)
            if len(bound_shapes) <= 1:
                text = "\n".join(
                    shape.text for shape in bound_shapes if getattr(shape, "has_text_frame", False)
                )
                substantial = [line for line in text.splitlines() if line.strip()]
                if len(substantial) >= 4:
                    warnings.append(
                        _warning(
                            "COMPOSITION.SUCCESS_CRITERIA_BLOCK",
                            record,
                            "Success criteria are rendered as one dense text block instead of separate checklist rows.",
                        )
                    )
            for shape in bound_shapes:
                lines = [line for line in shape.text.splitlines() if line.strip()] if getattr(shape, "has_text_frame", False) else []
                if len(lines) >= 6:
                    warnings.append(
                        _warning(
                            "COMPOSITION.DENSE_LIST",
                            record,
                            "A dense six-or-more-item list is contained in one text shape.",
                        )
                    )
                    break

        if role == "answer":
            font_groups = []
            for binding in bindings:
                shape = shapes.get(binding.get("shape_id"))
                if shape is not None:
                    sizes = _font_sizes(shape)
                    if sizes:
                        font_groups.append(max(sizes))
            if len(font_groups) >= 2 and max(font_groups) - min(font_groups) < 4:
                warnings.append(
                    _warning(
                        "COMPOSITION.FLAT_ANSWER_HIERARCHY",
                        record,
                        "Answer and explanation use effectively equal type emphasis.",
                    )
                )

        slide_area = presentation.slide_width * presentation.slide_height
        text_shapes = [shapes.get(binding.get("shape_id")) for binding in bindings]
        text_shapes = [shape for shape in text_shapes if shape is not None]
        for panel in slide.shapes:
            name = getattr(panel, "name", "") or ""
            if not name.startswith("DLP:panel:"):
                continue
            panel_area = panel.width * panel.height
            if panel_area / slide_area < 0.30:
                continue
            contained = [
                shape
                for shape in text_shapes
                if shape.left >= panel.left
                and shape.top >= panel.top
                and shape.left + shape.width <= panel.left + panel.width
                and shape.top + shape.height <= panel.top + panel.height
            ]
            if contained:
                occupancy = sum(shape.width * shape.height for shape in contained) / panel_area
                if occupancy < 0.18:
                    warnings.append(
                        _warning(
                            "COMPOSITION.LOW_OCCUPANCY_PANEL",
                            record,
                            f"Large panel contains only {occupancy:.0%} text-box occupancy.",
                        )
                    )
                    break

    for index in range(len(geometry) - 2):
        chunk = geometry[index : index + 3]
        signatures = [row[0] for row in chunk]
        roles = {row[1] for row in chunk}
        layouts = {row[2] for row in chunk}
        if signatures[0] and signatures[0] == signatures[1] == signatures[2] and (len(roles) > 1 or len(layouts) > 1):
            warnings.append(
                _warning(
                    "COMPOSITION.REPEATED_GEOMETRY",
                    records[index],
                    f"Slides {chunk[0][3]}-{chunk[2][3]} repeat the same major geometry despite differing semantic roles/layouts.",
                )
            )
            break

    status = "PASS" if not errors and not warnings else "FAIL"
    return {
        "schema_version": 1,
        "status": status,
        "artifact_sha256": sha256(deck_path),
        "render_manifest_sha256": sha256(render_manifest_path),
        "errors": errors,
        "warnings": warnings,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--deck", type=Path, required=True)
    parser.add_argument("--render-manifest", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    try:
        report = audit_visual_composition(args.deck, args.render_manifest)
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        report = {"schema_version": 1, "status": "FAIL", "errors": [str(exc)], "warnings": []}
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, indent=2))
    return 0 if report.get("status") == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
