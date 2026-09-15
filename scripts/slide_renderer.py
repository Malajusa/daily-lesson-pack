from __future__ import annotations

from dataclasses import asdict, dataclass, field
import hashlib
from pathlib import Path
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_VERTICAL_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from slide_render_model import BlockKind, ContentBlock, RenderPack, SlideRole, SlideSpec

FONT = "Trebuchet MS"
SLIDE_W = 13.333
SLIDE_H = 7.5
RAIL_W = 0.18
LEFT = 0.72
TOP_CONTENT = 1.50

ROLE_STYLE = {
    SlideRole.REMINDER: {"panel": "FFF3BF", "accent": "D6A900"},
    SlideRole.MODEL: {"panel": "FFF3BF", "accent": "D6A900"},
    SlideRole.WORKED_EXAMPLE: {"panel": "FFF3BF", "accent": "D6A900"},
    SlideRole.QUESTION: {"panel": "EAF2F8", "accent": "005A9C"},
    SlideRole.READING: {"panel": "EAF2F8", "accent": "005A9C"},
    SlideRole.MORNING_WORK: {"panel": "EAF2F8", "accent": "005A9C"},
    SlideRole.ANSWER: {"panel": "EAF7EE", "accent": "1B7F3A"},
    SlideRole.SUCCESS_CRITERIA: {"panel": "EAF7EE", "accent": "1B7F3A"},
    SlideRole.TRANSITION: {"panel": "FFFFFF", "accent": "103A5E"},
}


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass
class SlideRenderRecord:
    slide_number: int
    slide_id: str
    instance_id: str
    role: str
    layout_id: str
    source_blocks: list[str]
    shape_bindings: list[dict]
    groups: dict[str, dict] = field(default_factory=dict)
    metadata_labels: list[str] = field(default_factory=list)


@dataclass
class RenderManifest:
    schema_version: int
    renderer: str
    renderer_version: str
    slides: list[SlideRenderRecord]
    deck_sha256: str | None = None

    def to_dict(self) -> dict:
        return {
            "schema_version": self.schema_version,
            "renderer": self.renderer,
            "renderer_version": self.renderer_version,
            "deck_sha256": self.deck_sha256,
            "slides": [asdict(slide) for slide in self.slides],
        }


def estimate_text_lines(text: str, width_inches: float, font_pt: int) -> int:
    if width_inches <= 0:
        return 999
    chars_per_line = max(8, int(width_inches * (96.0 / max(font_pt, 1))))
    lines = 0
    for raw in text.splitlines() or [text]:
        words = raw.split()
        if not words:
            lines += 1
            continue
        current = 0
        count = 1
        for word in words:
            needed = len(word) + (1 if current else 0)
            if current + needed > chars_per_line:
                count += 1
                current = len(word)
            else:
                current += needed
        lines += count
    return max(1, lines)


def choose_font_size(text: str, role: SlideRole, width_inches: float, height_inches: float) -> int:
    floors = {
        SlideRole.READING: 28,
        SlideRole.MODEL: 28,
        SlideRole.SUCCESS_CRITERIA: 28,
        SlideRole.WORKED_EXAMPLE: 30,
        SlideRole.MORNING_WORK: 26,
        SlideRole.QUESTION: 36,
        SlideRole.ANSWER: 36,
        SlideRole.REMINDER: 30,
        SlideRole.TRANSITION: 30,
    }
    floor = floors[role]
    if role is SlideRole.ANSWER and len(text.strip()) <= 18:
        candidates = [54, 52, 48, 44, 40, 36]
    elif role is SlideRole.QUESTION and len(text.strip()) <= 60:
        candidates = [46, 44, 42, 40, 38, 36]
    elif role in {SlideRole.READING, SlideRole.MODEL}:
        candidates = [32, 30, 28]
    elif role is SlideRole.MORNING_WORK:
        candidates = [32, 30, 28, 26]
    elif role is SlideRole.SUCCESS_CRITERIA:
        candidates = [32, 30, 28]
    else:
        candidates = list(range(max(floor, 40), floor - 1, -2))
    for pt in candidates:
        lines = estimate_text_lines(text, width_inches, pt)
        if lines * (pt / 72) * 1.24 <= height_inches:
            return max(pt, floor)
    return floor


def panel_height_for_lines(lines: int, font_pt: int, *, min_height: float, max_height: float) -> float:
    return min(max_height, max(min_height, lines * (font_pt / 72.0) * 1.28 + 0.48))


def split_paragraph_for_projection(text: str, *, max_chars: int = 520) -> tuple[str, ...]:
    text = " ".join(text.split())
    if len(text) <= max_chars:
        return (text,)
    sentences = re.split(r"(?<=[.!?])\s+", text)
    chunks: list[str] = []
    current = ""
    for sentence in sentences:
        if not sentence:
            continue
        trial = f"{current} {sentence}".strip()
        if current and len(trial) > max_chars:
            chunks.append(current)
            current = sentence
        elif len(sentence) > max_chars:
            for word in sentence.split():
                trial_word = f"{current} {word}".strip()
                if current and len(trial_word) > max_chars:
                    chunks.append(current)
                    current = word
                else:
                    current = trial_word
        else:
            current = trial
    if current:
        chunks.append(current)
    return tuple(chunks)


def _add_rail(slide, spec: SlideSpec):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(RAIL_W), Inches(SLIDE_H))
    shape.name = f"DLP:rail:{spec.id}"
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(ROLE_STYLE[spec.role]["accent"])
    shape.line.fill.background()
    return shape


def _add_text(slide, text: str, x: float, y: float, w: float, h: float, *, font_pt: int, bold: bool = False,
              color: str = "17324D", align=PP_ALIGN.LEFT, valign=MSO_VERTICAL_ANCHOR.TOP,
              name: str | None = None, margins=(0.06, 0.06, 0.04, 0.04)):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = Inches(margins[0])
    frame.margin_right = Inches(margins[1])
    frame.margin_top = Inches(margins[2])
    frame.margin_bottom = Inches(margins[3])
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.name = FONT
        run.font.size = Pt(font_pt)
        run.font.bold = bold
        run.font.color.rgb = _rgb(color)
    return shape


def _add_panel(slide, x: float, y: float, w: float, h: float, *, fill_hex: str, accent_hex: str | None = None,
               name: str | None = None, radius: bool = True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill_hex)
    shape.line.color.rgb = _rgb(accent_hex or fill_hex)
    shape.line.width = Pt(1.2)
    return shape


def _header(slide, spec: SlideSpec, record: SlideRenderRecord):
    accent = ROLE_STYLE[spec.role]["accent"]
    _add_text(slide, spec.eyebrow, LEFT, 0.35, 7.7, 0.33, font_pt=16, bold=True, color=accent,
              name=f"DLP:meta:eyebrow:{spec.id}")
    _add_text(slide, spec.title, LEFT, 0.69, 11.5, 0.62, font_pt=28, bold=True,
              name=f"DLP:meta:title:{spec.id}")
    record.metadata_labels.extend([spec.eyebrow, spec.title])


def _bind(record: SlideRenderRecord, shape, block: ContentBlock, *, split_index: int | None = None):
    source_id = block.id if split_index is None else f"{block.id}#part{split_index}"
    record.shape_bindings.append({
        "source_block_id": source_id,
        "source_task_id": block.source_task_id,
        "source_field": block.source_field,
        "shape_id": shape.shape_id,
        "shape_name": shape.name,
    })
    if source_id not in record.source_blocks:
        record.source_blocks.append(source_id)


def _new_slide(prs: Presentation, spec: SlideSpec, slide_number: int, layout_id: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rail(slide, spec)
    record = SlideRenderRecord(slide_number, spec.id, spec.instance_id, spec.role.value, layout_id, [], [])
    _header(slide, spec, record)
    return slide, record


def _render_question(prs, spec: SlideSpec, slide_number: int):
    slide, record = _new_slide(prs, spec, slide_number, "question-primary-v1")
    main = spec.blocks[0]
    width = 11.45
    pt = choose_font_size(main.text, SlideRole.QUESTION, width - 0.5, 3.9)
    height = panel_height_for_lines(estimate_text_lines(main.text, width - 0.5, pt), pt, min_height=1.65, max_height=3.25)
    _add_panel(slide, LEFT, TOP_CONTENT, width, height, fill_hex=ROLE_STYLE[spec.role]["panel"],
               accent_hex=ROLE_STYLE[spec.role]["accent"], name=f"DLP:panel:{spec.id}:primary")
    shape = _add_text(slide, main.text, LEFT + 0.28, TOP_CONTENT + 0.22, width - 0.56, height - 0.44,
                      font_pt=pt, bold=True, name=f"DLP:block:{main.id}")
    _bind(record, shape, main)
    y = TOP_CONTENT + height + 0.25
    for block in spec.blocks[1:]:
        height2 = 1.0
        _add_panel(slide, LEFT, y, width, height2, fill_hex="FFFFFF", accent_hex="CBD5E1",
                   name=f"DLP:panel:{spec.id}:support")
        shape2 = _add_text(slide, block.text, LEFT + 0.25, y + 0.16, width - 0.5, height2 - 0.3,
                           font_pt=30, name=f"DLP:block:{block.id}")
        _bind(record, shape2, block)
        y += height2 + 0.18
    return [(slide, record)]


def _render_answer(prs, spec: SlideSpec, slide_number: int):
    slide, record = _new_slide(prs, spec, slide_number, "answer-primary-v1")
    main = spec.blocks[0]
    width = 11.45
    pt = choose_font_size(main.text, SlideRole.ANSWER, width - 0.6, 2.6)
    height = panel_height_for_lines(estimate_text_lines(main.text, width - 0.6, pt), pt, min_height=1.8, max_height=2.8)
    _add_panel(slide, LEFT, TOP_CONTENT, width, height, fill_hex=ROLE_STYLE[spec.role]["panel"],
               accent_hex=ROLE_STYLE[spec.role]["accent"], name=f"DLP:panel:{spec.id}:answer")
    shape = _add_text(slide, main.text, LEFT + 0.3, TOP_CONTENT + 0.2, width - 0.6, height - 0.4,
                      font_pt=pt, bold=True, color="123B22", name=f"DLP:block:{main.id}")
    _bind(record, shape, main)
    y = TOP_CONTENT + height + 0.28
    for block in spec.blocks[1:]:
        support_pt = min(34, max(28, pt - 12))
        height2 = panel_height_for_lines(estimate_text_lines(block.text, width - 0.5, support_pt), support_pt,
                                         min_height=1.0, max_height=1.65)
        _add_panel(slide, LEFT, y, width, height2, fill_hex="FFFFFF", accent_hex="B8D8C1",
                   name=f"DLP:panel:{spec.id}:explanation")
        shape2 = _add_text(slide, block.text, LEFT + 0.25, y + 0.16, width - 0.5, height2 - 0.3,
                           font_pt=support_pt, name=f"DLP:block:{block.id}")
        _bind(record, shape2, block)
        y += height2 + 0.18
    return [(slide, record)]


def _reading_parts(spec: SlideSpec):
    paragraph = next((block for block in spec.blocks if block.kind is BlockKind.PARAGRAPH), None)
    others = [block for block in spec.blocks if block is not paragraph]
    if paragraph is None:
        return [(None, others)]
    chunks = split_paragraph_for_projection(paragraph.text, max_chars=500)
    return [
        (ContentBlock(paragraph.id, paragraph.source_task_id, paragraph.source_field, paragraph.kind,
                      paragraph.visibility, chunk, paragraph.group), others if index == len(chunks) else [])
        for index, chunk in enumerate(chunks, 1)
    ]


def _render_reading(prs, spec: SlideSpec, slide_number: int):
    outputs = []
    parts = _reading_parts(spec)
    for part_index, (paragraph, others) in enumerate(parts, 1):
        slide, record = _new_slide(prs, spec, slide_number + len(outputs), "reading-measure-v1")
        record.slide_id = spec.id if part_index == 1 else f"{spec.id}-part{part_index}"
        measure = 8.0
        pt = 30 if paragraph and len(paragraph.text) < 360 else 28
        height = panel_height_for_lines(estimate_text_lines(paragraph.text if paragraph else "", measure - 0.5, pt), pt,
                                        min_height=2.15, max_height=4.15)
        _add_panel(slide, LEFT, TOP_CONTENT, measure, height, fill_hex=ROLE_STYLE[spec.role]["panel"],
                   accent_hex=ROLE_STYLE[spec.role]["accent"], name=f"DLP:panel:{spec.id}:reading")
        if paragraph:
            shape = _add_text(slide, paragraph.text, LEFT + 0.26, TOP_CONTENT + 0.18, measure - 0.52, height - 0.36,
                              font_pt=pt, name=f"DLP:block:{paragraph.id}:part{part_index}")
            _bind(record, shape, paragraph, split_index=(part_index if len(parts) > 1 else None))
        y = TOP_CONTENT + height + 0.20
        for block in others:
            qpt = 32
            qheight = panel_height_for_lines(estimate_text_lines(block.text, 10.9, qpt), qpt, min_height=1.05, max_height=1.55)
            _add_panel(slide, LEFT, y, 11.45, qheight, fill_hex="FFFFFF", accent_hex="9CC4DE",
                       name=f"DLP:panel:{spec.id}:question")
            shape2 = _add_text(slide, block.text, LEFT + 0.25, y + 0.14, 10.95, qheight - 0.26,
                               font_pt=qpt, bold=True, name=f"DLP:block:{block.id}")
            _bind(record, shape2, block)
            y += qheight + 0.15
        outputs.append((slide, record))
    return outputs


def _render_morning_work(prs, spec: SlideSpec, slide_number: int):
    slide, record = _new_slide(prs, spec, slide_number, "morning-work-split-v1")
    grouped: dict[str, list[ContentBlock]] = {}
    for block in spec.blocks:
        grouped.setdefault(block.group or "maths", []).append(block)
    ordered = [name for name in ("maths", "literacy", "extension") if name in grouped]
    if len(ordered) >= 2:
        col_width = 5.55
        gap = 0.28
        positions = {
            ordered[0]: (LEFT, TOP_CONTENT, col_width, 5.25),
            ordered[1]: (LEFT + col_width + gap, TOP_CONTENT, col_width, 5.25),
        }
        if "extension" in ordered[2:]:
            grouped[ordered[1]].extend(grouped["extension"])
    else:
        positions = {ordered[0]: (LEFT, TOP_CONTENT, 11.38, 5.25)}
    for group, (x, y, width, height) in positions.items():
        accent = "005A9C" if group == "maths" else "7A4F9D"
        fill = "EAF2F8" if group == "maths" else "F3ECF8"
        _add_panel(slide, x, y, width, height, fill_hex=fill, accent_hex=accent, name=f"DLP:group:{group}:{spec.id}")
        _add_text(slide, group.upper(), x + 0.22, y + 0.18, width - 0.44, 0.35, font_pt=18, bold=True,
                  color=accent, name=f"DLP:group-label:{group}:{spec.id}")
        blocks = grouped[group]
        each = max(0.62, min(1.05, (height - 0.72) / max(1, len(blocks))))
        current_y = y + 0.65
        for block in blocks:
            pt = choose_font_size(block.text, SlideRole.MORNING_WORK, width - 0.5, each - 0.08)
            shape = _add_text(slide, block.text, x + 0.25, current_y, width - 0.5, each - 0.08,
                              font_pt=pt, name=f"DLP:block:{block.id}")
            _bind(record, shape, block)
            current_y += each
        record.groups[group] = {"left_inches": x, "top_inches": y, "width_inches": width, "height_inches": height}
    return [(slide, record)]


def _render_success(prs, spec: SlideSpec, slide_number: int):
    slide, record = _new_slide(prs, spec, slide_number, "success-checklist-v1")
    count = len(spec.blocks)
    width = 11.45
    gap = 0.12
    row_height = min(0.82, max(0.58, (5.35 - gap * (count - 1)) / max(count, 1)))
    y = TOP_CONTENT
    for index, block in enumerate(spec.blocks, 1):
        _add_panel(slide, LEFT, y, width, row_height, fill_hex="EAF7EE", accent_hex="B8D8C1",
                   name=f"DLP:check-panel:{index}:{spec.id}")
        _add_text(slide, "☐", LEFT + 0.18, y + 0.08, 0.45, row_height - 0.12, font_pt=28, bold=True,
                  color="1B7F3A", name=f"DLP:check:{index}:{spec.id}")
        pt = choose_font_size(block.text, SlideRole.SUCCESS_CRITERIA, width - 0.95, row_height - 0.08)
        shape = _add_text(slide, block.text, LEFT + 0.72, y + 0.08, width - 0.94, row_height - 0.12,
                          font_pt=pt, name=f"DLP:block:{block.id}")
        _bind(record, shape, block)
        y += row_height + gap
    return [(slide, record)]


def _render_model_like(prs, spec: SlideSpec, slide_number: int, layout_id: str):
    slide, record = _new_slide(prs, spec, slide_number, layout_id)
    width = 9.15
    text = "\n\n".join(block.text for block in spec.blocks)
    pt = choose_font_size(text, spec.role, width - 0.55, 4.8)
    height = panel_height_for_lines(estimate_text_lines(text, width - 0.55, pt), pt, min_height=2.0, max_height=4.9)
    _add_panel(slide, LEFT, TOP_CONTENT, width, height, fill_hex=ROLE_STYLE[spec.role]["panel"],
               accent_hex=ROLE_STYLE[spec.role]["accent"], name=f"DLP:panel:{spec.id}:model")
    current_y = TOP_CONTENT + 0.20
    each = (height - 0.4) / max(1, len(spec.blocks))
    for block in spec.blocks:
        shape = _add_text(slide, block.text, LEFT + 0.28, current_y, width - 0.56, each - 0.10, font_pt=pt,
                          bold=(block.kind is BlockKind.WORKED_STEP), name=f"DLP:block:{block.id}")
        _bind(record, shape, block)
        current_y += each
    return [(slide, record)]


def _render_transition(prs, spec: SlideSpec, slide_number: int):
    slide, record = _new_slide(prs, spec, slide_number, "transition-v1")
    block = spec.blocks[0]
    shape = _add_text(slide, block.text, 1.25, 2.7, 10.8, 1.3, font_pt=42, bold=True,
                      align=PP_ALIGN.CENTER, valign=MSO_VERTICAL_ANCHOR.MIDDLE, name=f"DLP:block:{block.id}")
    _bind(record, shape, block)
    return [(slide, record)]


def _render_spec(prs, spec: SlideSpec, slide_number: int):
    if spec.role is SlideRole.QUESTION:
        return _render_question(prs, spec, slide_number)
    if spec.role is SlideRole.ANSWER:
        return _render_answer(prs, spec, slide_number)
    if spec.role is SlideRole.READING:
        return _render_reading(prs, spec, slide_number)
    if spec.role is SlideRole.MORNING_WORK:
        return _render_morning_work(prs, spec, slide_number)
    if spec.role is SlideRole.SUCCESS_CRITERIA:
        return _render_success(prs, spec, slide_number)
    if spec.role is SlideRole.WORKED_EXAMPLE:
        return _render_model_like(prs, spec, slide_number, "worked-example-v1")
    if spec.role is SlideRole.MODEL:
        return _render_model_like(prs, spec, slide_number, "model-paragraph-v1")
    if spec.role is SlideRole.REMINDER:
        return _render_model_like(prs, spec, slide_number, "reminder-primary-v1")
    if spec.role is SlideRole.TRANSITION:
        return _render_transition(prs, spec, slide_number)
    raise ValueError(f"unsupported slide role {spec.role}")


def render_presentation(pack: RenderPack, out_path: Path) -> RenderManifest:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    presentation = Presentation()
    presentation.slide_width = Inches(SLIDE_W)
    presentation.slide_height = Inches(SLIDE_H)
    records: list[SlideRenderRecord] = []
    slide_number = 1
    for spec in pack.slides:
        rendered = _render_spec(presentation, spec, slide_number)
        records.extend(record for _, record in rendered)
        slide_number += len(rendered)
    presentation.save(out_path)
    return RenderManifest(1, "daily-lesson-pack", "1.0.0", records, _sha256(out_path))
