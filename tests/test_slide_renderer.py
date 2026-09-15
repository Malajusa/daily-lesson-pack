import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from slide_render_model import BlockKind, ComponentKind, ContentBlock, ProjectionPolicy, RenderPack, SlideRole, SlideSpec, Visibility
from slide_renderer import render_presentation


def block(identifier, text, kind=BlockKind.TASK, field="prompt", group=None):
    return ContentBlock(identifier, identifier.split(":")[0], field, kind, Visibility.STUDENT, text, group)


def spec(role, blocks, component=ComponentKind.MATHEMATICS):
    titles = {
        SlideRole.QUESTION: "Your turn", SlideRole.ANSWER: "Answer", SlideRole.READING: "Read and think",
        SlideRole.MORNING_WORK: "Morning Work", SlideRole.SUCCESS_CRITERIA: "Success criteria",
    }
    return SlideSpec(
        "s-" + role.value, "i", component, role,
        "MORNING WORK" if component is ComponentKind.MORNING_WORK else "MATHEMATICS",
        titles[role], tuple(blocks), ProjectionPolicy(28, 70, 10, True), role.value,
    )


def shape_by_source(slide, record, source):
    binding = next(item for item in record.shape_bindings if item["source_block_id"] == source)
    return next(shape for shape in slide.shapes if shape.shape_id == binding["shape_id"])


def font_size(shape):
    return max(run.font.size.pt for paragraph in shape.text_frame.paragraphs for run in paragraph.runs if run.font.size)


class SlideRendererTests(unittest.TestCase):
    def render(self, slide_spec):
        temp = tempfile.TemporaryDirectory()
        path = Path(temp.name) / "deck.pptx"
        manifest = render_presentation(RenderPack("run", "", (slide_spec,)), path)
        return temp, path, manifest

    def test_question_visual_grammar(self):
        temp, path, manifest = self.render(spec(SlideRole.QUESTION, [block("t:prompt", "Calculate 125 × 8.")]))
        try:
            presentation = Presentation(path)
            slide = presentation.slides[0]
            record = manifest.slides[0]
            self.assertEqual(record.layout_id, "question-primary-v1")
            rail = next(shape for shape in slide.shapes if shape.name.startswith("DLP:rail:"))
            self.assertEqual(str(rail.fill.fore_color.rgb), "005A9C")
            self.assertGreaterEqual(font_size(shape_by_source(slide, record, "t:prompt")), 36)
        finally:
            temp.cleanup()

    def test_answer_hierarchy(self):
        temp, path, manifest = self.render(spec(SlideRole.ANSWER, [
            block("t:answer", "24 cm³", BlockKind.ANSWER, "answer"),
            block("t:why_answer", "Because 3 × 4 × 2 = 24.", BlockKind.SUPPORT, "why_answer"),
        ]))
        try:
            slide = Presentation(path).slides[0]
            answer = font_size(shape_by_source(slide, manifest.slides[0], "t:answer"))
            explanation = font_size(shape_by_source(slide, manifest.slides[0], "t:why_answer"))
            self.assertGreaterEqual(answer, 48)
            self.assertLessEqual(explanation, answer - 6)
        finally:
            temp.cleanup()

    def test_reading_measure_and_fail_closed_split(self):
        temp, path, manifest = self.render(spec(SlideRole.READING, [
            block("t:paragraph", "The storm rolled across the harbour.", BlockKind.PARAGRAPH, "paragraph"),
            block("t:prompt", "What changed?", BlockKind.TASK, "prompt"),
        ], ComponentKind.SHARED_READING))
        try:
            self.assertLessEqual(shape_by_source(Presentation(path).slides[0], manifest.slides[0], "t:paragraph").width, Inches(8.2))
        finally:
            temp.cleanup()

        long_text = " ".join(["The valley changed after the storm and the river climbed over its banks."] * 24)
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaisesRegex(ValueError, "canonical split"):
                render_presentation(RenderPack("run", "", (spec(SlideRole.READING, [
                    block("t:paragraph", long_text, BlockKind.PARAGRAPH, "paragraph"),
                    block("t:prompt", "How?", BlockKind.TASK, "prompt"),
                ], ComponentKind.SHARED_READING),)), Path(folder) / "too-long.pptx")

    def test_morning_work_has_two_broad_groups_and_structural_labels(self):
        slide_spec = spec(SlideRole.MORNING_WORK, [
            block("m:prompt", "345 + 278", group="maths"),
            block("l:prompt", "Correct: we was late", group="literacy"),
        ], ComponentKind.MORNING_WORK)
        temp, _, manifest = self.render(slide_spec)
        try:
            record = manifest.slides[0]
            self.assertEqual(set(record.groups), {"maths", "literacy"})
            self.assertTrue(all(group["width_inches"] >= 5 for group in record.groups.values()))
            self.assertEqual({item["kind"] for item in record.structural_text}, {"eyebrow", "title", "group_label"})
        finally:
            temp.cleanup()

    def test_success_criteria_are_separate_rows(self):
        blocks = [block(f"t:success_criteria:{i}", text, BlockKind.CHECK, "success_criteria") for i, text in enumerate(["A", "B", "C", "D", "E", "F"], 1)]
        temp, path, _ = self.render(spec(SlideRole.SUCCESS_CRITERIA, blocks))
        try:
            self.assertEqual(len([shape for shape in Presentation(path).slides[0].shapes if shape.name.startswith("DLP:check:")]), 6)
        finally:
            temp.cleanup()


if __name__ == "__main__":
    unittest.main()
