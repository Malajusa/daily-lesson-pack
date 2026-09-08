from __future__ import annotations

import importlib.util
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]


def load_audit():
    path = ROOT / "scripts/audit_panel_containment.py"
    spec = importlib.util.spec_from_file_location("panel_audit_explicit", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def add_panel(slide, owner: str, left=1.0, top=1.0, width=5.0, height=3.0):
    panel = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    panel.name = f"DLP:panel:{owner}"
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(245, 245, 245)
    return panel


def add_text(slide, text: str, owner: str | None, left=1.3, top=1.3, width=4.3, height=1.0):
    shape = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.name = "DLP:main" + (f"|panel={owner}" if owner is not None else "")
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(36)
    return shape


class ExplicitPanelOwnershipTests(unittest.TestCase):
    def setUp(self) -> None:
        self.audit = load_audit()

    def audit_slide(self, slide):
        prs = slide.part.package.presentation_part.presentation
        return self.audit.audit_slide(slide, 1, prs.slide_width, prs.slide_height, 0.15)

    def test_explicit_owned_text_inside_panel_passes(self) -> None:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_panel(slide, "morning-literacy")
        add_text(slide, "Add the commas and write the sentence.", "morning-literacy")
        result = self.audit.audit_slide(slide, 1, prs.slide_width, prs.slide_height, 0.15)
        self.assertNotEqual(result["status"], "fail", result)

    def test_explicit_owned_text_crossing_panel_is_hard_failure(self) -> None:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_panel(slide, "morning-literacy", left=1.0, top=1.0, width=4.0, height=2.0)
        add_text(
            slide,
            "Extension text that deliberately crosses the panel edge.",
            "morning-literacy",
            left=1.3,
            top=2.4,
            width=4.2,
            height=1.0,
        )
        result = self.audit.audit_slide(slide, 1, prs.slide_width, prs.slide_height, 0.15)
        codes = {issue["code"] for issue in result["issues"]}
        self.assertIn("explicit_text_crosses_panel_bounds", codes, result)
        self.assertEqual(result["status"], "fail")

    def test_missing_explicit_panel_is_hard_failure(self) -> None:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_text(slide, "This text claims an owner that does not exist.", "missing-panel")
        result = self.audit.audit_slide(slide, 1, prs.slide_width, prs.slide_height, 0.15)
        codes = {issue["code"] for issue in result["issues"]}
        self.assertIn("explicit_panel_missing", codes, result)
        self.assertEqual(result["status"], "fail")

    def test_unowned_dlp_content_fails_when_slide_uses_explicit_panels(self) -> None:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_panel(slide, "morning-literacy")
        add_text(slide, "Owned task", "morning-literacy", top=1.3, height=0.7)
        add_text(slide, "Unowned extension", None, top=2.2, height=0.7)
        result = self.audit.audit_slide(slide, 1, prs.slide_width, prs.slide_height, 0.15)
        codes = {issue["code"] for issue in result["issues"]}
        self.assertIn("dlp_text_missing_panel_owner", codes, result)
        self.assertEqual(result["status"], "fail")


if __name__ == "__main__":
    unittest.main()
