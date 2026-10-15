"""Positive/negative structural regressions using synthetic PPTX fixtures."""
import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import sys
import tempfile
import unittest
from pptx import Presentation
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
import audit_pack_contract as audit
import literacy_reading_audit as modes


def slide(p,title,body='',meta=None,marker='DLP_LITERACY:'):
    s=p.slides.add_slide(p.slide_layouts[6])
    box=s.shapes.add_textbox(Inches(.5), Inches(.3), Inches(12), Inches(.7))
    box.text=title
    box.text_frame.paragraphs[0].font.size=Pt(24)
    box=s.shapes.add_textbox(Inches(.7), Inches(1.8), Inches(11.8), Inches(4.5))
    box.text=body
    for para in box.text_frame.paragraphs:para.font.size=Pt(36)
    if 'REMINDER' in title or 'SUPPORT' in title:
        from pptx.dml.color import RGBColor
        box.fill.solid();box.fill.fore_color.rgb=RGBColor.from_string('FFF3BF')
    if meta is not None:s.notes_slide.notes_text_frame.text=marker+json.dumps(meta,ensure_ascii=False)
    return s


def literacy_deck(check=True,retry=False,wrong_label=False,unknown=False):
    p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5)
    words=['unsafe','unhelpful']
    for i,word in enumerate(words,1):
        mode='question_first_check' if check and i==2 else 'supported_practice'
        roles=['QUESTION','SUPPORT','RETRY','ANSWER'] if mode=='question_first_check' and retry else ['QUESTION','ANSWER'] if mode=='question_first_check' else ['REMINDER','QUESTION','ANSWER']
        target='suffix-less' if unknown and i==2 else 'prefix-un'
        for role in roles:
            data=dict(target_rule=target,mode=mode,selection_source='synthetic teacher fixture',history_known=not unknown,prior_cues=[])
            data['cue_status']='supported' if role in {'REMINDER','SUPPORT','RETRY'} or mode=='supported_practice' else 'unknown' if unknown else 'recently_cued'
            if wrong_label and role=='QUESTION' and i==2:data['cue_status']='no_identified_recent_target_cue'
            content=f"Write the meaning of ‘{word}’." if role in {'QUESTION','RETRY'} else "Remember: The prefix ‘un-’ can mean ‘not’." if role in {'REMINDER','SUPPORT'} else f"{word.capitalize()} means not {word[2:]}."
            slide(p,f'LITERACY WARM-UP {i} OF 2 — {role}',content,data)
    return p


def reading_deck(folder,mutation=None):
    from reportlab.pdfgen import canvas
    source='Maya hid the sign behind the pots.\n\nLater she raised the sign. People read it at the entrance.'
    path=Path(folder)/'source.pdf';c=canvas.Canvas(str(path));c.drawString(40,780,'Maya hid the sign behind the pots.');c.drawString(40,750,'Later she raised the sign. People read it at the entrance.');c.save()
    p=Presentation()
    common=dict(source_id='stall',mode='whole_text')
    rows=[dict(common,role='SOURCE',source_text=source,handout_path='source.pdf',handout_sha256=hashlib.sha256(path.read_bytes()).hexdigest()),dict(common,role='MODEL',think_aloud='I connect the hidden sign with people reading it later.'),dict(common,role='QUESTION',stage='guided',selection_source='fixture',evidence_quotes=['hid the sign','raised the sign']),dict(common,role='ANSWER',stage='guided',evidence_quotes=['hid the sign','raised the sign']),dict(common,role='QUESTION',stage='reduced_support',selection_source='fixture',evidence_quotes=['People read it at the entrance.']),dict(common,role='ANSWER',stage='reduced_support',evidence_quotes=['People read it at the entrance.'])]
    if mutation:mutation(rows)
    for i,data in enumerate(rows):
        title=f'SHARED READING {i//2} OF 2 — {data["role"]}' if data['role'] in {'QUESTION','ANSWER'} else f'SHARED READING — {data["role"]}'
        slide(p,title,'Source available on the complete handout.',data,'DLP_READING:')
    return p


class LiteracyReadingTests(unittest.TestCase):
    def check(self,p,count=2):
        with tempfile.TemporaryDirectory() as d:
            path=Path(d)/'deck.pptx';p.save(path)
            return audit.audit_deck(str(path),literacy_count=count,numeracy_count=0)[0]
    def test_legacy_default_ten_supported_items(self):
        p=Presentation()
        for n in range(1,11):
            for role in ['REMINDER','QUESTION','ANSWER']:
                slide(p,f'LITERACY WARM-UP {n} OF 10 — {role}','Remember: A prefix changes meaning.' if role=='REMINDER' else 'Word meaning')
        self.assertEqual(self.check(p,10),[]);self.assertEqual(len(p.slides),30)
    def test_selected_check_after_same_rule_feedback(self):self.assertEqual(self.check(literacy_deck()),[])
    def test_separate_retry_path(self):self.assertEqual(self.check(literacy_deck(retry=True)),[])
    def test_unknown_exposure(self):self.assertEqual(self.check(literacy_deck(unknown=True)),[])
    def test_false_uncued_label_after_feedback_rejected(self):self.assertIn('literacy_cue_label_inaccurate',{x['code'] for x in self.check(literacy_deck(wrong_label=True))})
    def test_unknown_cannot_be_uncued(self):self.assertIn('literacy_cue_label_inaccurate',{x['code'] for x in self.check(literacy_deck(unknown=True,wrong_label=True))})
    def test_intervening_content_does_not_clear_cue(self):
        self.assertEqual(modes.cue_status('QUESTION','prefix-un',{'prefix-un','conjunction'},[],True),'recently_cued')
    def test_known_cue_with_unknown_history(self):self.assertEqual(modes.cue_status('QUESTION','un',{'un'},[],False),'recently_cued')
    def test_missing_selection_rejected(self):
        p=literacy_deck();s=p.slides[3];d=modes.record(s,'DLP_LITERACY:');d.pop('selection_source');s.notes_slide.notes_text_frame.text='DLP_LITERACY:'+json.dumps(d)
        self.assertIn('literacy_selection_missing',{x['code'] for x in self.check(p)})
    def test_invalid_selection_types_rejected(self):
        for value in [None,False,[],{}]:
            with self.subTest(value=value):
                p=literacy_deck();sl=p.slides[3];d=modes.record(sl,'DLP_LITERACY:');d['selection_source']=value;sl.notes_slide.notes_text_frame.text='DLP_LITERACY:'+json.dumps(d)
                self.assertIn('literacy_selection_missing',{x['code'] for x in self.check(p)})
    def test_overwritten_attempt_rejected(self):
        p=literacy_deck(retry=True);s=p.slides[3];d=modes.record(s,'DLP_LITERACY:');d['response']='correct';s.notes_slide.notes_text_frame.text='DLP_LITERACY:'+json.dumps(d)
        self.assertIn('literacy_attempts_not_separated',{x['code'] for x in self.check(p)})
    def test_reading_default_pairs(self):
        p=Presentation()
        for role in ['QUESTION','ANSWER']:slide(p,f'SHARED READING 1 OF 1 — {role}','Maya raised the sign.')
        self.assertEqual(self.check(p,0),[])
    def read_check(self,mutation=None):
        with tempfile.TemporaryDirectory() as d:
            p=reading_deck(d,mutation);path=Path(d)/'deck.pptx';p.save(path)
            return audit.audit_deck(str(path),literacy_count=0,numeracy_count=0)[0]
    def test_accessible_whole_text(self):self.assertEqual(self.read_check(),[])
    def test_missing_handout(self):self.assertIn('reading_source_access_unverified',{x['code'] for x in self.read_check(lambda r:r[0].update(handout_path='absent.pdf'))})
    def test_changed_handout(self):self.assertIn('reading_handout_hash_mismatch',{x['code'] for x in self.read_check(lambda r:r[0].update(handout_sha256='bad'))})
    def test_unsupported_source_quote(self):self.assertIn('reading_evidence_quote_unsupported',{x['code'] for x in self.read_check(lambda r:r[3].update(evidence_quotes=['They bought the tools.']))})
    def test_missing_think_aloud(self):self.assertIn('reading_think_aloud_missing',{x['code'] for x in self.read_check(lambda r:r[1].update(think_aloud=''))})
    def test_invalid_think_aloud_types_rejected(self):
        for value in [None,False,[],{}]:
            with self.subTest(value=value):
                self.assertIn('reading_think_aloud_missing',{x['code'] for x in self.read_check(lambda r:r[1].update(think_aloud=value))})
    def test_false_transfer(self):self.assertIn('reading_false_transfer_claim',{x['code'] for x in self.read_check(lambda r:r[4].update(independent_transfer=True))})
    def test_reading_invalid_selection_types_rejected(self):
        for value in [None,False,[],{}]:
            with self.subTest(value=value):
                self.assertIn('reading_selection_missing',{x['code'] for x in self.read_check(lambda r:r[2].update(selection_source=value))})
    def test_mismatched_answer_stage_rejected(self):
        self.assertIn('reading_feedback_identity_mismatch',{x['code'] for x in self.read_check(lambda r:r[3].update(stage='reduced_support'))})
    def test_malformed_source_type_rejected(self):
        self.assertIn('reading_source_empty',{x['code'] for x in self.read_check(lambda r:r[0].update(source_text=[]))})
    def test_package_dependency_closure(self):
        from build_chatgpt_package import build_file_map
        _,files=build_file_map(ROOT)
        for name in ['scripts/literacy_reading_audit.py','references/literacy-reading-modes.md','skills/dlp-pack-qa/scripts/literacy_reading_audit.py','skills/dlp-literacy-warmup/references/literacy-reading-modes.md']:
            self.assertIn(name,files)

if __name__=='__main__':unittest.main()
