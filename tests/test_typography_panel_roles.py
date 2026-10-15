"""The renderer's ownership suffix must preserve the declared text-role floor."""
import sys
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from audit_slide_typography import audit_slide


class PanelRoleTypographyTests(unittest.TestCase):
    def font_failures(self, role, size):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(3))
        box.name = role
        box.text = 'The two partial products combine to give the same whole.'
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(size)
        result = audit_slide(slide, 1, prs.slide_width, prs.slide_height, True)
        return [issue for issue in result['issues'] if issue['code'] == 'undersized_body_text']

    def test_owned_support_retains_28_point_floor(self):
        for role in ['why', 'instruction', 'explanation', 'cue']:
            with self.subTest(role=role):
                self.assertEqual(self.font_failures('DLP:' + role + '|panel=answer', 28), [])

    def test_owned_support_below_floor_still_fails(self):
        self.assertTrue(self.font_failures('DLP:why|panel=answer', 27))

    def test_main_and_unknown_roles_cannot_claim_support_floor(self):
        for role in ['main', 'why-not']:
            with self.subTest(role=role):
                self.assertTrue(self.font_failures('DLP:' + role + '|panel=answer', 28))


if __name__ == '__main__':
    unittest.main()
