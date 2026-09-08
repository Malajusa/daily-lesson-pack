from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "scripts" / "build_daily_pack.py"


class BuildDailyPackCliTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.source = self.root / "source"
        self.source.mkdir()
        self.run_root = self.root / "run"

        def field(value: object, source: str = "current-run fixture") -> dict:
            return {"resolved": True, "value": value, "source": source}

        self.request = {
            "schema_version": 1,
            "requested_date": "2026-09-08",
            "requested_output": "powerpoint",
        }
        self.context = {
            "schema_version": 2,
            "active_year_profile": field("year-4-5", "year-level profile"),
            "pack_profile": field("default-daily-pack-architecture"),
            "date": field("2026-09-08"),
            "term_week": field("T3W8"),
            "day": field("Tuesday"),
            "timetable": field("08:30 Morning Work"),
            "mathematics_focus": field("mass"),
            "english_focus": field("narrative"),
            "lesson_status": field("previous scheduled lesson covered"),
            "timetable_instances": [
                {
                    "id": "morning-work-1",
                    "owner": "dlp-morning-work",
                    "start": "08:30",
                    "duration_minutes": 10,
                    "purpose": "independent retrieval",
                }
            ],
            "required_artifacts": ["deck"],
        }
        self.components = {
            "schema_version": 2,
            "generation_run_id": "generation-cli-1",
            "active_year_profile": "year-4-5",
            "instances": [
                {
                    "instance_id": "morning-work-1",
                    "owner": "dlp-morning-work",
                    "status": "PASS",
                    "active_year_profile": "year-4-5",
                    "estimated_minutes": 9,
                    "slide_range": [1, 1],
                    "checks": [
                        {
                            "id": "COMPONENT.CONTENT",
                            "result": "PASS",
                            "evidence": "The generated task is tied to the resolved instance.",
                        }
                    ],
                }
            ],
        }
        self.content = {
            "schema_version": 3,
            "instances": self.context["timetable_instances"],
            "tasks": [
                {
                    "id": "morning-q1",
                    "instance_id": "morning-work-1",
                    "operation": "retrieve",
                    "fields": {"prompt": "4 386 + 2 749 =", "answer": "7 135"},
                    "demands": [
                        {"id": "response", "action": "calculate", "answer_quote": "7 135"}
                    ],
                }
            ],
            "documents": {},
        }

        self.request_path = self.source / "request.json"
        self.context_path = self.source / "context.json"
        self.content_path = self.source / "content.json"
        self.component_path = self.source / "component-record.json"
        self.deck_path = self.source / "pack.pptx"
        self.manifest_path = self.source / "manifest.json"
        self.request_path.write_text(json.dumps(self.request), encoding="utf-8")
        self.context_path.write_text(json.dumps(self.context), encoding="utf-8")
        self.content_path.write_text(json.dumps(self.content), encoding="utf-8")
        self.component_path.write_text(json.dumps(self.components), encoding="utf-8")
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(self.deck_path)
        self.manifest_path.write_text(
            json.dumps({"schema_version": 3, "artifacts": [], "bindings": []}),
            encoding="utf-8",
        )

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def command(self, *extra: str) -> list[str]:
        return [
            sys.executable,
            str(BUILD),
            "--request",
            str(self.request_path),
            "--context",
            str(self.context_path),
            "--content",
            str(self.content_path),
            "--component-record",
            str(self.component_path),
            "--deck",
            str(self.deck_path),
            "--manifest",
            str(self.manifest_path),
            "--out",
            str(self.run_root),
            *extra,
        ]

    def run_build(self, *extra: str) -> subprocess.CompletedProcess[str]:
        return subprocess.run(
            self.command(*extra),
            cwd=ROOT,
            text=True,
            capture_output=True,
            check=False,
        )

    def test_cli_stages_candidate_without_creating_release(self) -> None:
        result = self.run_build()
        self.assertEqual(result.returncode, 0, result.stderr)
        payload = json.loads(result.stdout)
        self.assertEqual(payload["status"], "CANDIDATE")
        self.assertTrue((self.run_root / "candidate" / "pack.pptx").is_file())
        self.assertFalse((self.run_root / "released").exists())

    def test_cli_invalid_context_fails_before_candidate_staging(self) -> None:
        self.context["mathematics_focus"]["resolved"] = False
        self.context_path.write_text(json.dumps(self.context), encoding="utf-8")
        result = self.run_build()
        self.assertEqual(result.returncode, 1)
        payload = json.loads(result.stdout)
        self.assertEqual(payload["status"], "FAIL")
        self.assertIn("mathematics_focus", payload["reason"])
        self.assertFalse((self.run_root / "candidate").exists())
        self.assertFalse((self.run_root / "released").exists())

    def test_cli_incomplete_release_evidence_cannot_promote_candidate(self) -> None:
        evidence_dir = self.root / "evidence"
        evidence_dir.mkdir()
        result = self.run_build("--release-evidence-dir", str(evidence_dir))
        self.assertEqual(result.returncode, 0, result.stderr)
        payload = json.loads(result.stdout)
        self.assertEqual(payload["status"], "CANDIDATE")
        self.assertIn("Release evidence is incomplete", payload["reason"])
        self.assertTrue((self.run_root / "candidate" / "pack.pptx").is_file())
        self.assertFalse((self.run_root / "released" / "pack.pptx").exists())


if __name__ == "__main__":
    unittest.main()
