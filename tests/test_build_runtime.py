from __future__ import annotations

import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from dlp_build_runtime import (  # noqa: E402
    promote_release,
    sha256,
    stage_candidate,
    validate_component_record,
    validate_resolved_context,
)


class BuildRuntimeTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.source = self.root / "source"
        self.source.mkdir()
        self.run_root = self.root / "build" / "run-1"

        def field(value: object, source: str = "current-run fixture") -> dict:
            return {"resolved": True, "value": value, "source": source}

        self.context = {
            "schema_version": 2,
            "active_year_profile": field("year-4-5", "year-level profile"),
            "pack_profile": field("default-daily-pack-architecture"),
            "date": field("2026-09-08"),
            "term_week": field("T3W8"),
            "day": field("Tuesday"),
            "timetable": field("08:30 Morning Work"),
            "mathematics_focus": field("mass"),
            "english_focus": field("narrative"),
            "lesson_status": field("previous scheduled lesson covered"),
            "timetable_instances": [
                {
                    "id": "morning-work-1",
                    "owner": "dlp-morning-work",
                    "start": "08:30",
                    "duration_minutes": 10,
                    "purpose": "independent retrieval",
                }
            ],
            "required_artifacts": ["deck"],
        }
        self.components = {
            "schema_version": 2,
            "generation_run_id": "generation-1",
            "active_year_profile": "year-4-5",
            "instances": [
                {
                    "instance_id": "morning-work-1",
                    "owner": "dlp-morning-work",
                    "status": "PASS",
                    "active_year_profile": "year-4-5",
                    "estimated_minutes": 9,
                    "slide_range": [1, 1],
                    "checks": [
                        {
                            "id": "COMPONENT.CONTENT",
                            "result": "PASS",
                            "evidence": "The generated task is tied to the resolved instance.",
                        }
                    ],
                }
            ],
        }
        self.request = {
            "schema_version": 1,
            "requested_date": "2026-09-08",
            "requested_output": "powerpoint",
        }
        self.content = {
            "schema_version": 3,
            "instances": self.context["timetable_instances"],
            "tasks": [
                {
                    "id": "morning-q1",
                    "instance_id": "morning-work-1",
                    "operation": "retrieve",
                    "fields": {"prompt": "4 386 + 2 749 =", "answer": "7 135"},
                    "demands": [
                        {"id": "response", "action": "calculate", "answer_quote": "7 135"}
                    ],
                }
            ],
            "documents": {},
        }

        self.request_path = self.source / "request.json"
        self.context_path = self.source / "context.json"
        self.content_path = self.source / "content.json"
        self.component_path = self.source / "component-record.json"
        self.deck_path = self.source / "pack.pptx"
        self.manifest_path = self.source / "manifest.json"
        self.request_path.write_text(json.dumps(self.request), encoding="utf-8")
        self.context_path.write_text(json.dumps(self.context), encoding="utf-8")
        self.content_path.write_text(json.dumps(self.content), encoding="utf-8")
        self.component_path.write_text(json.dumps(self.components), encoding="utf-8")
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(self.deck_path)
        self.manifest_path.write_text(
            json.dumps({"schema_version": 3, "artifacts": [], "bindings": []}),
            encoding="utf-8",
        )

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_valid_resolved_context_passes(self) -> None:
        self.assertEqual(validate_resolved_context(self.context), [])

    def test_unresolved_context_blocks_build(self) -> None:
        self.context["mathematics_focus"]["resolved"] = False
        self.assertTrue(any("mathematics_focus" in e for e in validate_resolved_context(self.context)))

    def test_missing_scheduled_component_blocks_build(self) -> None:
        self.components["instances"] = []
        errors = validate_component_record(self.context, self.components)
        self.assertTrue(any("morning-work-1" in e for e in errors))

    def test_component_fail_blocks_build(self) -> None:
        self.components["instances"][0]["status"] = "FAIL"
        errors = validate_component_record(self.context, self.components)
        self.assertTrue(any("PASS" in e for e in errors))

    def test_component_profile_mismatch_blocks_build(self) -> None:
        self.components["instances"][0]["active_year_profile"] = "year-6"
        errors = validate_component_record(self.context, self.components)
        self.assertTrue(any("year profile" in e.lower() for e in errors))

    def test_stage_candidate_never_creates_released_output(self) -> None:
        result = stage_candidate(
            self.run_root,
            request_path=self.request_path,
            context_path=self.context_path,
            content_path=self.content_path,
            component_record_path=self.component_path,
            deck_path=self.deck_path,
            manifest_path=self.manifest_path,
        )
        self.assertEqual(result["status"], "CANDIDATE")
        self.assertTrue((self.run_root / "candidate" / "pack.pptx").is_file())
        self.assertFalse((self.run_root / "released").exists())

    def test_hand_written_pass_without_release_evidence_cannot_promote(self) -> None:
        stage_candidate(
            self.run_root,
            request_path=self.request_path,
            context_path=self.context_path,
            content_path=self.content_path,
            component_record_path=self.component_path,
            deck_path=self.deck_path,
            manifest_path=self.manifest_path,
        )
        fake = self.run_root / "candidate" / "release.json"
        fake.write_text(json.dumps({"status": "PASS"}), encoding="utf-8")
        result = promote_release(self.run_root / "candidate", self.run_root / "released", {})
        self.assertEqual(result["status"], "CANDIDATE")
        self.assertFalse((self.run_root / "released" / "pack.pptx").exists())

    def test_successful_release_requires_audit_output_bound_to_current_deck(self) -> None:
        stage_candidate(
            self.run_root,
            request_path=self.request_path,
            context_path=self.context_path,
            content_path=self.content_path,
            component_record_path=self.component_path,
            deck_path=self.deck_path,
            manifest_path=self.manifest_path,
        )
        candidate = self.run_root / "candidate"
        evidence = {}
        for key in (
            "contract",
            "year_profile",
            "typography",
            "containment",
            "visual",
            "semantic_review",
            "warning_ledger",
            "visual_review",
            "semantic_trace",
            "visual_trace",
        ):
            path = candidate / f"{key}.json"
            path.write_text("{}", encoding="utf-8")
            evidence[key] = path

        def runner(command, **kwargs):
            out = Path(command[command.index("--out") + 1])
            out.write_text(
                json.dumps(
                    {
                        "status": "PASS",
                        "artifact_sha256": sha256(candidate / "pack.pptx"),
                        "manifest_sha256": sha256(candidate / "manifest.json"),
                        "failures": [],
                    }
                ),
                encoding="utf-8",
            )
            return type("Result", (), {"returncode": 0, "stderr": "", "stdout": ""})()

        result = promote_release(candidate, self.run_root / "released", evidence, runner=runner)
        self.assertEqual(result["status"], "RELEASED")
        self.assertTrue((self.run_root / "released" / "pack.pptx").is_file())

    def test_deck_mutation_after_audit_prevents_promotion(self) -> None:
        stage_candidate(
            self.run_root,
            request_path=self.request_path,
            context_path=self.context_path,
            content_path=self.content_path,
            component_record_path=self.component_path,
            deck_path=self.deck_path,
            manifest_path=self.manifest_path,
        )
        candidate = self.run_root / "candidate"
        evidence = {}
        for key in (
            "contract",
            "year_profile",
            "typography",
            "containment",
            "visual",
            "semantic_review",
            "warning_ledger",
            "visual_review",
            "semantic_trace",
            "visual_trace",
        ):
            path = candidate / f"{key}.json"
            path.write_text("{}", encoding="utf-8")
            evidence[key] = path

        def runner(command, **kwargs):
            out = Path(command[command.index("--out") + 1])
            old_hash = sha256(candidate / "pack.pptx")
            out.write_text(
                json.dumps(
                    {
                        "status": "PASS",
                        "artifact_sha256": old_hash,
                        "manifest_sha256": sha256(candidate / "manifest.json"),
                        "failures": [],
                    }
                ),
                encoding="utf-8",
            )
            with (candidate / "pack.pptx").open("ab") as stream:
                stream.write(b"changed-after-review")
            return type("Result", (), {"returncode": 0, "stderr": "", "stdout": ""})()

        result = promote_release(candidate, self.run_root / "released", evidence, runner=runner)
        self.assertEqual(result["status"], "CANDIDATE")
        self.assertFalse((self.run_root / "released" / "pack.pptx").exists())


if __name__ == "__main__":
    unittest.main()
