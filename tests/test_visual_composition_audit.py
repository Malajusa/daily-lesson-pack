import hashlib
import json
import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
from audit_visual_composition import audit_visual_composition


def add_text(slide, text, x, y, width, height, size=28, bold=False, name="DLP:block:test"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    shape.name = name
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = text
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.name = "Trebuchet MS"
    return shape


class CompositionAuditTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.p = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def manifest(self, deck, slides):
        path = self.p / "render.json"
        path.write_text(json.dumps({
            "schema_version": 1,
            "renderer": "daily-lesson-pack",
            "renderer_version": "test",
            "deck_sha256": hashlib.sha256(deck.read_bytes()).hexdigest(),
            "slides": slides,
        }))
        return path

    def test_flags_dense_success_criteria(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = add_text(slide, "Clear opening\nPrecise verbs\nVaried sentences\nCorrect punctuation\nParagraphs\nStrong ending", 1, 1.5, 11, 4.5, 32, True)
        deck = self.p / "dense.pptx"
        presentation.save(deck)
        record = {
            "slide_number": 1, "slide_id": "s", "instance_id": "i", "role": "success_criteria", "layout_id": "success-checklist-v1",
            "source_blocks": ["t:success_criteria"],
            "shape_bindings": [{"source_block_id": "t:success_criteria", "source_task_id": "t", "source_field": "success_criteria", "shape_id": shape.shape_id, "shape_name": shape.name}],
            "groups": {}, "metadata_labels": [],
        }
        codes = {warning["code"] for warning in audit_visual_composition(deck, self.manifest(deck, [record]))["warnings"]}
        self.assertIn("COMPOSITION.SUCCESS_CRITERIA_BLOCK", codes)

    def test_flags_overwide_paragraph_and_repeated_metadata(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        paragraph = add_text(slide, "A paragraph across the classroom surface.", .6, 1.5, 11.8, 2, 28, name="DLP:block:p")
        add_text(slide, "Tuesday 15 September 2026 | Term 3, Week 9", .6, 6.7, 8, .3, 14, name="DLP:meta:date")
        deck = self.p / "wide.pptx"
        presentation.save(deck)
        record = {
            "slide_number": 2, "slide_id": "s", "instance_id": "i", "role": "reading", "layout_id": "reading-measure-v1",
            "source_blocks": ["t:paragraph"],
            "shape_bindings": [{"source_block_id": "t:paragraph", "source_task_id": "t", "source_field": "paragraph", "shape_id": paragraph.shape_id, "shape_name": paragraph.name}],
            "groups": {}, "metadata_labels": ["Tuesday 15 September 2026 | Term 3, Week 9"],
        }
        codes = {warning["code"] for warning in audit_visual_composition(deck, self.manifest(deck, [record]))["warnings"]}
        self.assertTrue({"COMPOSITION.PARAGRAPH_MEASURE", "COMPOSITION.REPEATED_METADATA"} <= codes)

    def test_flags_teacher_content_and_missing_morning_group(self):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        teacher = add_text(slide, "Ask two students to justify.", 1, 5.5, 8, .6, 18, name="DLP:block:teacher")
        deck = self.p / "teacher.pptx"
        presentation.save(deck)
        record = {
            "slide_number": 1, "slide_id": "s", "instance_id": "i", "role": "morning_work", "layout_id": "morning-work-split-v1",
            "source_blocks": ["t:teacher_note"],
            "shape_bindings": [{"source_block_id": "t:teacher_note", "source_task_id": "t", "source_field": "teacher_note", "shape_id": teacher.shape_id, "shape_name": teacher.name}],
            "groups": {}, "metadata_labels": [],
        }
        codes = {warning["code"] for warning in audit_visual_composition(deck, self.manifest(deck, [record]))["warnings"]}
        self.assertIn("COMPOSITION.TEACHER_CONTENT", codes)
        self.assertIn("COMPOSITION.MORNING_WORK_GROUPING", codes)


if __name__ == "__main__":
    unittest.main()
