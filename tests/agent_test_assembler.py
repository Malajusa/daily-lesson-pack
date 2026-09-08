"""Synthetic canonical-binding fixture. NEVER a production lesson renderer.

PNG files explicitly contain synthetic labels, not rendered lesson evidence.
This helper only tests staging/hash/coverage plumbing; release tests must fail.
"""
import hashlib
import json
import sys
from pathlib import Path
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    request=json.load(sys.stdin)
    root=Path(request['output_directory'])
    content=json.loads((root/'content.json').read_text())
    artifacts=[];bindings=[];renders=[]
    for role,name in [('deck','pack.pptx'),('briefing','briefing.pptx')]:
        prs=Presentation()
        for task in content['tasks']:
            for field in ('prompt','answer'):
                slide=prs.slides.add_slide(prs.slide_layouts[6])
                shape=slide.shapes.add_textbox(Inches(1),Inches(1),Inches(8),Inches(4))
                shape.text=task['fields'][field]
                page=len(prs.slides)
                bindings.append(dict(artifact=role,record=task['id'],field=field,page=page,shape_id=shape.shape_id))
        prs.save(root/name)
        artifacts.append(dict(id=role,role=role,path=name,pages=len(prs.slides),sha256=digest(root/name)))
        for n in range(1,len(prs.slides)+1):
            path=f'renders/{role}-{n}.png'
            (root/path).parent.mkdir(exist_ok=True)
            image=Image.new('RGB',(300,200),'white')
            ImageDraw.Draw(image).text((10,10),'SYNTHETIC TEST / NOT RENDERED',fill='black')
            image.save(root/path)
            renders.append(dict(artifact=role,page=n,path=path,sha256=digest(root/path),artifact_sha256=digest(root/name)))
    manifest=dict(schema_version=3,content_sha256=digest(root/'content.json'),context_sha256=digest(root/'context.json'),
                  artifacts=artifacts,bindings=bindings,renders=renders)
    (root/'manifest.json').write_text(json.dumps(manifest))
    print(json.dumps({'status':'PASS'}))


if __name__=='__main__':
    main()
