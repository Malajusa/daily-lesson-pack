"""Exercise real review decisions with clearly synthetic host receipts."""
import json
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import pack_evidence as evidence


class ReleaseBoundaryWaiverTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name)
        self.repo = self.path / "repo"
        (self.repo / "references").mkdir(parents=True)
        self.write(self.repo / "references/qa-requirements.json", {
            "requirements": [{"id": "CENTRAL.OPTIONAL", "owner": "dlp-maths-lesson",
                              "scope": "instance", "method": "semantic"}]})
        self.root_patch = patch.object(evidence, "ROOT", self.repo)
        self.root_patch.start()
        self.addCleanup(self.root_patch.stop)
        self.content = {
            "schema_version": 3,
            "instances": [{"id": "maths-1", "owner": "dlp-maths-lesson"}],
            "tasks": [{"id": "task-1", "instance_id": "maths-1", "operation": "partition",
                       "fields": {"prompt": "Split the outline.", "answer": "Two rectangles."},
                       "demands": [{"id": "partition", "action": "partition",
                                    "answer_quote": "Two rectangles."}]}]}
        self.set_boundary("dlp-maths-lesson", "MATHS.ADVANCE.ROUTE")
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
        shape.text = "Split the outline."
        deck.save(self.path / "deck.pptx")
        self.manifest = {
            "schema_version": 3,
            "artifacts": [{"id": "deck", "path": "deck.pptx", "role": "deck",
                           "pages": 1, "sha256": evidence.digest(self.path / "deck.pptx")}],
            "bindings": [{"artifact": "deck", "record": "task-1", "field": "prompt",
                          "page": 1, "shape_id": shape.shape_id}]}
        self.citation = {"artifact": "deck", "page": 1, "shape_id": shape.shape_id,
                         "quote": "Split the outline."}
        self.write(self.path / "content.json", self.content)
        self.write(self.path / "manifest.json", self.manifest)
        (self.path / "synthetic-transcript.txt").write_text("SYNTHETIC TEST FIXTURE ONLY")

    @staticmethod
    def write(path, value):
        path.write_text(json.dumps(value), encoding="utf-8")

    def set_boundary(self, owner, boundary_id):
        path = self.repo / "skills" / owner / "references/release-boundaries.json"
        path.parent.mkdir(parents=True, exist_ok=True)
        self.write(path, {"owner": owner, "release_policy": "fail_closed",
                          "boundaries": [{"id": boundary_id, "statement": "Required fixture check."}]})
        return path

    def make_review(self):
        self.write(self.path / "content.json", self.content)
        self.review = {
            "manifest_sha256": evidence.digest(self.path / "manifest.json"),
            "content_sha256": evidence.digest(self.path / "content.json"),
            "requirements_sha256": evidence.requirements_digest(self.content),
            "execution_id": "synthetic-review",
            "checks": [{"id": key, "subject": subject, "result": "PASS",
                        "observation": "Synthetic fixture observation.",
                        "citations": [dict(self.citation)]}
                       for key, subject in sorted(evidence.expected_checks(
                           self.content, self.manifest, "semantic"))]}
        return self.review

    def audit(self):
        self.write(self.path / "review.json", self.review)
        self.write(self.path / "synthetic-trace.json", {
            "source": "external-runner", "execution_id": "synthetic-review",
            "reviewer_actor": "synthetic-reviewer", "generator_actor": "synthetic-author",
            "review_sha256": evidence.digest(self.path / "review.json"),
            "manifest_sha256": evidence.digest(self.path / "manifest.json"),
            "transcript_path": "synthetic-transcript.txt",
            "transcript_sha256": evidence.digest(self.path / "synthetic-transcript.txt")})
        return evidence.audit_review(
            self.path / "review.json", "semantic", self.path / "manifest.json",
            self.path / "content.json", self.content, self.manifest,
            self.path / "synthetic-trace.json")

    def check(self, key):
        return next(item for item in self.review["checks"] if item["id"] == key)

    def test_registered_non_numeracy_boundary_rejects_na(self):
        for boundary_id in ("MATHS.ADVANCE.ROUTE", "CUSTOM.REQUIRED"):
            with self.subTest(boundary_id=boundary_id):
                self.set_boundary("dlp-maths-lesson", boundary_id)
                self.make_review()
                self.check(boundary_id).update(
                    result="NA", applicability_reason="Synthetic attempted waiver.")
                self.assertTrue(any("Release boundary cannot be waived" in error
                                    for error in self.audit()))

    def test_registered_boundary_fail_and_omission_rejected(self):
        for result in ("FAIL", "OMITTED"):
            with self.subTest(result=result):
                self.make_review()
                check = self.check("MATHS.ADVANCE.ROUTE")
                if result == "OMITTED":
                    self.review["checks"].remove(check)
                else:
                    check["result"] = result
                self.assertTrue(self.audit())

    def test_registered_boundary_pass_accepted(self):
        self.make_review()
        self.assertEqual(self.audit(), [])

    def test_numeracy_na_remains_rejected(self):
        self.content["instances"][0]["owner"] = "dlp-numeracy-warmup"
        self.set_boundary("dlp-numeracy-warmup", "NW.CALIBRATION")
        self.make_review()
        self.check("NW.CALIBRATION").update(result="NA", applicability_reason="Fixture.")
        self.assertTrue(any("Release boundary cannot be waived" in e for e in self.audit()))

    def test_unrelated_central_na_remains_allowed(self):
        self.make_review()
        self.check("CENTRAL.OPTIONAL").update(
            result="NA", applicability_reason="Applicable synthetic central exemption.")
        self.assertEqual(self.audit(), [])

    def test_unscheduled_owner_adds_no_checks(self):
        self.set_boundary("dlp-numeracy-warmup", "NW.CALIBRATION")
        checks = evidence.expected_checks(self.content, self.manifest, "semantic")
        self.assertNotIn(("NW.CALIBRATION", "maths-1"), checks)

    def test_new_pedagogical_fail_blocks_review(self):
        self.make_review()
        self.check("MATHS.ADVANCE.ROUTE")["result"] = "FAIL"
        self.assertTrue(any("incomplete/failed check" in error and
                            "MATHS.ADVANCE.ROUTE" in error for error in self.audit()))

    def test_missing_new_pedagogical_check_blocks_review(self):
        self.make_review()
        self.review["checks"].remove(self.check("MATHS.ADVANCE.ROUTE"))
        self.assertTrue(any("missing check" in error and "MATHS.ADVANCE.ROUTE" in error
                            for error in self.audit()))

    def test_changed_boundary_invalidates_old_requirements_digest(self):
        self.make_review()
        path = self.repo / "skills/dlp-maths-lesson/references/release-boundaries.json"
        contract = json.loads(path.read_text())
        contract["boundaries"][0]["statement"] = "A stronger required fixture check."
        self.write(path, contract)
        self.assertIn("semantic: review uses stale requirements", self.audit())


if __name__ == "__main__":
    unittest.main()
