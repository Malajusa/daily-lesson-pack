from __future__ import annotations

import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
VALIDATOR = ROOT / "skills/dlp-numeracy-warmup/scripts/validate_warmup_deck.py"


def add_text(slide, text, x, y, w=2.6, h=0.55, size=36, name=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text = text
    if name:
        shape.name = name
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.name = "Trebuchet MS"
    return shape


def make_pair(prs, number, *, why_on_question=False, labels=("ALL", "MOST", "SOME"), graph=False, sparse_scale=False):
    q = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(q, f"NUMERACY WARM-UP {number} OF 5 • QUESTION", 0.4, 0.25, 8.5, 0.5, 24)
    add_text(q, "Reading a column graph" if graph else "Addition", 0.4, 0.9, 8.5, 0.5, 28)
    for idx, label in enumerate(labels):
        x = 0.7 + idx * 3.0
        add_text(q, label, x, 2.0, 2.3, 0.45, 22)
        add_text(q, f"{10*number + idx} + {idx + 2} = ___", x, 2.55, 2.3, 0.7, 36, "DLP:main")
    if graph:
        ticks = [0, 5, 10] if sparse_scale else [0, 1, 2, 3]
        for idx, tick in enumerate(ticks):
            add_text(q, str(tick), 0.2, 3.8 + idx * 0.35, 0.5, 0.3, 16)
    if why_on_question:
        add_text(q, "WHY", 0.7, 5.7, 1.0, 0.4, 22)
        add_text(q, "Why does this work?", 1.8, 5.7, 6.5, 0.5, 28, "DLP:why")

    a = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(a, f"NUMERACY WARM-UP {number} OF 5 • ANSWER", 0.4, 0.25, 8.5, 0.5, 24)
    add_text(a, "Reading a column graph: answers" if graph else "Addition: answers", 0.4, 0.9, 8.5, 0.5, 28)
    for idx, label in enumerate(labels):
        x = 0.7 + idx * 3.0
        add_text(a, label, x, 2.0, 2.3, 0.45, 22)
        add_text(a, str(10*number + idx + idx + 2), x, 2.55, 2.3, 0.7, 36, "DLP:main")
    add_text(a, "WHY", 0.7, 5.7, 1.0, 0.4, 22)
    add_text(a, "The place-value relationship stays consistent.", 1.8, 5.7, 6.5, 0.5, 28, "DLP:why")


def write_deck(path: Path, *, grouped=False, why_on_question=False, labels=("ALL", "MOST", "SOME"), graph_pair=None, sparse_scale=False):
    prs = Presentation()
    if not grouped:
        for number in range(1, 6):
            make_pair(prs, number, why_on_question=why_on_question, labels=labels,
                      graph=(number == graph_pair), sparse_scale=sparse_scale)
    else:
        for number in range(1, 6):
            q = prs.slides.add_slide(prs.slide_layouts[6])
            add_text(q, f"NUMERACY WARM-UP {number} OF 5 • QUESTION", 0.4, 0.25, 8.5, 0.5, 24)
            for idx, label in enumerate(labels):
                x = 0.7 + idx * 3.0
                add_text(q, label, x, 2.0, 2.3, 0.45, 22)
                add_text(q, f"{number + idx} + 2 = ___", x, 2.55, 2.3, 0.7, 36, "DLP:main")
        for number in range(1, 6):
            a = prs.slides.add_slide(prs.slide_layouts[6])
            add_text(a, f"NUMERACY WARM-UP {number} OF 5 • ANSWER", 0.4, 0.25, 8.5, 0.5, 24)
            for idx, label in enumerate(labels):
                x = 0.7 + idx * 3.0
                add_text(a, label, x, 2.0, 2.3, 0.45, 22)
                add_text(a, str(number + idx + 2), x, 2.55, 2.3, 0.7, 36, "DLP:main")
            add_text(a, "WHY", 0.7, 5.7, 1.0, 0.4, 22)
            add_text(a, "Reasoning.", 1.8, 5.7, 6.5, 0.5, 28, "DLP:why")
    prs.save(path)


class StandaloneSkillContractTests(unittest.TestCase):
    def test_skill_is_standalone_first_and_preserves_exact_slide_grammar(self):
        text = (ROOT / "skills/dlp-numeracy-warmup/SKILL.md").read_text(encoding="utf-8")
        self.assertIn("Standalone mode", text)
        self.assertIn("QUESTION 1 → ANSWER 1", text)
        self.assertIn("literal labels `ALL`, `MOST`, `SOME`", text)
        self.assertIn("Pairs 4 and 5", text)
        self.assertIn("curriculum alignment is unverified", text)
        self.assertNotIn("Before generation, read `references/year-level-context-contract.md`", text)

    def test_agent_default_prompt_carries_nonnegotiable_shape(self):
        text = (ROOT / "skills/dlp-numeracy-warmup/agents/openai.yaml").read_text(encoding="utf-8")
        for phrase in ("standalone mode", "QUESTION", "ANSWER", "ALL", "MOST", "SOME", "WHY", "validate"):
            self.assertIn(phrase, text)

    def test_version_is_bumped(self):
        self.assertEqual((ROOT / "VERSION").read_text().strip(), "3.10.0-rc.2")
        provenance = json.loads((ROOT / "RELEASE-PROVENANCE.json").read_text())
        self.assertEqual(provenance["version"], "3.10.0-rc.2")

    def test_validator_exists(self):
        self.assertTrue(VALIDATOR.is_file())


class WarmupDeckValidatorTests(unittest.TestCase):
    def run_validator(self, deck: Path):
        return subprocess.run(
            [sys.executable, str(VALIDATOR), str(deck), "--expected-pairs", "5"],
            capture_output=True, text=True, check=False,
        )

    def test_valid_canonical_deck_passes(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "valid.pptx"
            write_deck(deck, graph_pair=4, sparse_scale=False)
            result = self.run_validator(deck)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_grouped_questions_then_answers_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "grouped.pptx"
            write_deck(deck, grouped=True)
            result = self.run_validator(deck)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("adjacent", (result.stdout + result.stderr).lower())

    def test_numbered_tiers_fail(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "numbered.pptx"
            write_deck(deck, labels=("1", "2", "3"))
            result = self.run_validator(deck)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("all", (result.stdout + result.stderr).lower())

    def test_why_on_question_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "why-question.pptx"
            write_deck(deck, why_on_question=True)
            result = self.run_validator(deck)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("why", (result.stdout + result.stderr).lower())

    def test_sparse_graph_scale_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            deck = Path(tmp) / "sparse-graph.pptx"
            write_deck(deck, graph_pair=4, sparse_scale=True)
            result = self.run_validator(deck)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("scale", (result.stdout + result.stderr).lower())

    def test_existing_bad_output_fails(self):
        deck = Path("/mnt/data/review/year_2_3_numeracy_warmup_graphs_review.pptx")
        if not deck.is_file():
            self.skipTest("review fixture unavailable")
        result = self.run_validator(deck)
        self.assertNotEqual(result.returncode, 0)


class StandalonePackagingTests(unittest.TestCase):
    def test_component_package_contains_validator_and_excludes_private_settings_runtime(self):
        with tempfile.TemporaryDirectory() as tmp:
            result = subprocess.run(
                [sys.executable, str(ROOT / "scripts/package_component_skills.py"), "--out", tmp],
                capture_output=True, text=True, check=False,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            package = Path(tmp) / "dlp-numeracy-warmup.zip"
            with zipfile.ZipFile(package) as archive:
                names = set(archive.namelist())
            prefix = "dlp-numeracy-warmup/"
            self.assertIn(prefix + "scripts/validate_warmup_deck.py", names)
            self.assertNotIn(prefix + "scripts/teacher_context_store.py", names)
            self.assertNotIn(prefix + "references/creator-settings-contract.md", names)
            self.assertNotIn(prefix + "config/creator-defaults.json", names)

    def test_component_package_passes_dependency_audit_without_profile_registry(self):
        with tempfile.TemporaryDirectory() as tmp:
            build = Path(tmp) / "build"
            install = Path(tmp) / "install"
            result = subprocess.run(
                [sys.executable, str(ROOT / "scripts/package_component_skills.py"), "--out", str(build)],
                capture_output=True, text=True, check=False,
            )
            self.assertEqual(result.returncode, 0, result.stderr)
            package = build / "dlp-numeracy-warmup.zip"
            with zipfile.ZipFile(package) as archive:
                archive.extractall(install)
            audit = subprocess.run(
                [sys.executable, str(ROOT / "scripts/audit_package_dependencies.py"),
                 "--skill-root", str(install / "dlp-numeracy-warmup"), "--component"],
                capture_output=True, text=True, check=False,
            )
            self.assertEqual(audit.returncode, 0, audit.stdout + audit.stderr)


if __name__ == "__main__":
    unittest.main()
