import sys
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))


class RepoRenderedPackAuditTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.p = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def make_deck(self, extra="Your turn"):
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        bound = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1))
        bound.text = "Calculate 2 + 3."
        bound.name = "DLP:block:t:prompt"
        title = slide.shapes.add_textbox(Inches(1), Inches(.2), Inches(8), Inches(.5))
        title.text = extra
        title.name = "DLP:meta:title:s1"
        path = self.p / "deck.pptx"
        presentation.save(path)
        return path, bound.shape_id, title.shape_id

    def test_sanitiser_blanks_only_validated_structural_text_and_preserves_shape_ids(self):
        from audit_repo_rendered_pack import sanitise_structural_text
        deck, bound_id, title_id = self.make_deck()
        render = {
            "schema_version": 1, "renderer": "daily-lesson-pack",
            "slides": [{"slide_number": 1, "structural_text": [{"shape_id": title_id, "shape_name": "DLP:meta:title:s1", "text": "Your turn", "kind": "title"}]}],
        }
        out = self.p / "sanitised.pptx"
        errors = sanitise_structural_text(deck, render, out, bound_shapes={(1, bound_id)})
        self.assertEqual(errors, [])
        slide = Presentation(out).slides[0]
        by_id = {shape.shape_id: shape for shape in slide.shapes}
        self.assertEqual(by_id[bound_id].text, "Calculate 2 + 3.")
        self.assertEqual(by_id[title_id].text, "")

    def test_sanitiser_rejects_unlisted_unbound_text(self):
        from audit_repo_rendered_pack import sanitise_structural_text
        deck, bound_id, _ = self.make_deck("Sneak in an extra instruction")
        render = {"schema_version": 1, "renderer": "daily-lesson-pack", "slides": [{"slide_number": 1, "structural_text": []}]}
        out = self.p / "sanitised.pptx"
        errors = sanitise_structural_text(deck, render, out, bound_shapes={(1, bound_id)})
        self.assertTrue(any("unrecognised unbound renderer text" in error.lower() for error in errors))


if __name__ == "__main__":
    unittest.main()
