#!/usr/bin/env python3
"""Fail-closed structural validator for standalone DLP Numeracy Warm-up decks."""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

HEADER_RE = re.compile(
    r"NUMERACY\s+WARM[- ]?UP\s+(\d+)\s+OF\s+(\d+).*?\b(QUESTION|ANSWER)\b",
    re.I | re.S,
)
GRAPH_RE = re.compile(r"\b(?:graph|pictograph|chart)\b", re.I)
KEY_RE = re.compile(r"(?:★|●|■|◆|▲|○|□|♦|✦|✶|✹)\s*=\s*\d+", re.I)
SCALE_RE = re.compile(
    r"\b(?:each|every|one|1)\s+(?:interval|square|step|tick|grid\s*line)\s*(?:is|=|represents)\s*\d+\b",
    re.I,
)
TIER_LABELS = ("ALL", "MOST", "SOME")
LABEL_TOLERANCE = Inches(0.30)


def shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return "\n".join(p.text for p in shape.text_frame.paragraphs).strip()


def slide_text(slide) -> str:
    return "\n".join(text for shape in slide.shapes if (text := shape_text(shape)))


def exact_text_shapes(slide, text: str):
    wanted = text.strip().upper()
    return [shape for shape in slide.shapes if shape_text(shape).strip().upper() == wanted]


def effective_sizes(shape) -> list[float]:
    sizes: list[float] = []
    if not getattr(shape, "has_text_frame", False):
        return sizes
    for paragraph in shape.text_frame.paragraphs:
        paragraph_size = getattr(paragraph.font, "size", None)
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            size = run.font.size or paragraph_size
            if size is not None:
                sizes.append(float(size.pt))
    return sizes


def closest_task_shape(slide, label_shape, label_shapes):
    centres = sorted((shape.left + shape.width / 2, shape) for shape in label_shapes)
    centre = label_shape.left + label_shape.width / 2
    idx = [shape for _, shape in centres].index(label_shape)
    left_edge = 0 if idx == 0 else (centres[idx - 1][0] + centre) / 2
    right_edge = 10**12 if idx == len(centres) - 1 else (centre + centres[idx + 1][0]) / 2
    bottom = label_shape.top + label_shape.height
    candidates = []
    for shape in slide.shapes:
        text = shape_text(shape).strip()
        if not text or shape in label_shapes:
            continue
        shape_centre = shape.left + shape.width / 2
        if not (left_edge <= shape_centre <= right_edge):
            continue
        if shape.top < bottom:
            continue
        if text.upper() in {"WHY", *TIER_LABELS}:
            continue
        candidates.append((shape.top - bottom, shape))
    return min(candidates, default=(None, None), key=lambda item: item[0])[1]


def graph_scale_is_exact(slide) -> bool:
    text = slide_text(slide)
    if KEY_RE.search(text) or SCALE_RE.search(text):
        return True
    ticks = set()
    for shape in slide.shapes:
        value = shape_text(shape).strip()
        if re.fullmatch(r"-?\d+", value):
            ticks.add(int(value))
    if len(ticks) < 3:
        return False
    ordered = sorted(ticks)
    for start in range(len(ordered) - 2):
        if ordered[start + 1] == ordered[start] + 1 and ordered[start + 2] == ordered[start] + 2:
            return True
    return False


def validate(path: Path, expected_pairs: int = 5) -> dict:
    issues: list[dict] = []
    try:
        prs = Presentation(path)
    except Exception as exc:
        return {"status": "FAIL", "issues": [{"code": "deck_unreadable", "detail": str(exc)}]}

    expected_slides = expected_pairs * 2
    if len(prs.slides) != expected_slides:
        issues.append({
            "code": "slide_count",
            "message": f"Expected {expected_slides} slides for {expected_pairs} adjacent question/answer pairs; found {len(prs.slides)}.",
        })

    parsed = []
    for slide_no, slide in enumerate(prs.slides, 1):
        full_text = slide_text(slide)
        match = HEADER_RE.search(full_text)
        if not match:
            issues.append({
                "code": "header_missing",
                "slide": slide_no,
                "message": "Slide lacks canonical NUMERACY WARM-UP n OF total QUESTION/ANSWER header.",
            })
            parsed.append(None)
            continue
        number, total, role = int(match.group(1)), int(match.group(2)), match.group(3).upper()
        parsed.append((number, total, role))
        if total != expected_pairs:
            issues.append({
                "code": "header_total",
                "slide": slide_no,
                "message": f"Header total must be {expected_pairs}; found {total}.",
            })

        labels = {label: exact_text_shapes(slide, label) for label in TIER_LABELS}
        for label in TIER_LABELS:
            if len(labels[label]) != 1:
                issues.append({
                    "code": "tier_label",
                    "slide": slide_no,
                    "message": f"Slide must contain literal label {label} exactly once.",
                })
        label_shapes = [labels[label][0] for label in TIER_LABELS if len(labels[label]) == 1]
        if len(label_shapes) == 3:
            ordered = [shape_text(shape).strip().upper() for shape in sorted(label_shapes, key=lambda s: s.left)]
            if ordered != list(TIER_LABELS):
                issues.append({
                    "code": "tier_order",
                    "slide": slide_no,
                    "message": "Tier labels must appear left-to-right as ALL, MOST, SOME.",
                })
            for label in TIER_LABELS:
                task = closest_task_shape(slide, labels[label][0], label_shapes)
                if task is None:
                    issues.append({
                        "code": "tier_task_missing",
                        "slide": slide_no,
                        "message": f"{label} has no task/answer beneath its label.",
                    })
                    continue
                sizes = effective_sizes(task)
                if not sizes or min(sizes) < 36:
                    issues.append({
                        "code": "tier_font_floor",
                        "slide": slide_no,
                        "message": f"{label} task/answer must be at least 36 pt.",
                    })

        why_shapes = exact_text_shapes(slide, "WHY")
        if role == "QUESTION":
            if why_shapes:
                issues.append({
                    "code": "why_on_question",
                    "slide": slide_no,
                    "message": "WHY/reasoning belongs on the answer slide only.",
                })
            if GRAPH_RE.search(full_text) and not graph_scale_is_exact(slide):
                issues.append({
                    "code": "graph_scale",
                    "slide": slide_no,
                    "message": "Graph scale/key is not explicit enough to recover exact numeric answers.",
                })
        else:
            if len(why_shapes) != 1:
                issues.append({
                    "code": "why_missing",
                    "slide": slide_no,
                    "message": "Answer slide must contain one literal WHY label.",
                })
            else:
                why = why_shapes[0]
                candidates = []
                for shape in slide.shapes:
                    text = shape_text(shape).strip()
                    if not text or shape == why or text.upper() in {"WHY", *TIER_LABELS}:
                        continue
                    if shape.left >= why.left + why.width and abs(shape.top - why.top) <= Inches(0.75):
                        candidates.append((shape.left - (why.left + why.width), shape))
                explanation = min(candidates, default=(None, None), key=lambda item: item[0])[1]
                if explanation is None:
                    issues.append({
                        "code": "why_explanation_missing",
                        "slide": slide_no,
                        "message": "Answer WHY panel needs a concise mathematical explanation.",
                    })
                else:
                    sizes = effective_sizes(explanation)
                    if not sizes or min(sizes) < 28:
                        issues.append({
                            "code": "why_font_floor",
                            "slide": slide_no,
                            "message": "WHY explanation must be at least 28 pt.",
                        })

    if len(parsed) >= 2:
        for number in range(1, expected_pairs + 1):
            q_index = (number - 1) * 2
            a_index = q_index + 1
            expected_q = (number, expected_pairs, "QUESTION")
            expected_a = (number, expected_pairs, "ANSWER")
            actual_q = parsed[q_index] if q_index < len(parsed) else None
            actual_a = parsed[a_index] if a_index < len(parsed) else None
            if actual_q != expected_q or actual_a != expected_a:
                issues.append({
                    "code": "pair_adjacent",
                    "pair": number,
                    "message": f"Pair {number} must be adjacent QUESTION then ANSWER; found {actual_q!r} then {actual_a!r}.",
                })
                continue
            q_slide = prs.slides[q_index]
            a_slide = prs.slides[a_index]
            for label in TIER_LABELS:
                q_labels = exact_text_shapes(q_slide, label)
                a_labels = exact_text_shapes(a_slide, label)
                if len(q_labels) == len(a_labels) == 1:
                    if abs(q_labels[0].left - a_labels[0].left) > LABEL_TOLERANCE or abs(q_labels[0].top - a_labels[0].top) > LABEL_TOLERANCE:
                        issues.append({
                            "code": "tier_mirror",
                            "pair": number,
                            "message": f"{label} label position must mirror between question and answer slides.",
                        })

    return {"status": "PASS" if not issues else "FAIL", "issues": issues, "slides": len(prs.slides)}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("deck", type=Path)
    parser.add_argument("--expected-pairs", type=int, default=5)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    report = validate(args.deck, args.expected_pairs)
    payload = json.dumps(report, indent=2)
    if args.out:
        args.out.write_text(payload + "\n", encoding="utf-8")
    print(payload)
    return 0 if report["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
